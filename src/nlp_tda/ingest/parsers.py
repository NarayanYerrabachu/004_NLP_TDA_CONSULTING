from __future__ import annotations

import csv
import hashlib
import io
import shutil
import zipfile
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from nlp_tda.config import settings

# Parsed by the pipeline when found as files on disk.
SUPPORTED_SUFFIXES = {
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".csv",
    ".txt",
    ".md",
    ".eml",
    ".msg",
}

# Accepted by the ingest UI / upload API (includes archives that expand first).
UPLOAD_SUFFIXES = SUPPORTED_SUFFIXES | {".zip"}

LATER_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".gif", ".webp", ".heic"}


@dataclass
class ParsedDocument:
    path: Path
    title: str
    text: str
    mime_hint: str
    content_hash: str
    pages: list[str] = field(default_factory=list)
    warning: str | None = None


@dataclass
class FileProbe:
    name: str
    suffix: str
    detected_type: str
    status: str  # supported | archive | later | unsupported
    message: str = ""


def content_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def probe_filename(name: str) -> FileProbe:
    suffix = Path(name).suffix.lower()
    if suffix in SUPPORTED_SUFFIXES:
        return FileProbe(name, suffix, suffix.lstrip(".") or "file", "supported")
    if suffix == ".zip":
        return FileProbe(name, suffix, "zip", "archive", "Will be expanded; nested files parsed individually.")
    if suffix in LATER_SUFFIXES:
        return FileProbe(
            name,
            suffix,
            "image",
            "later",
            "Image OCR is out of MVP scope. Export text or use a text-layer PDF.",
        )
    return FileProbe(name, suffix or "(none)", "unknown", "unsupported", f"Type {suffix or 'unknown'} not supported yet.")


def parse_file(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    raw = path.read_bytes()
    digest = content_hash(raw)
    if suffix == ".pdf":
        return _parse_pdf(path, digest)
    if suffix == ".docx":
        return _parse_docx(path, digest)
    if suffix == ".pptx":
        return _parse_pptx(path, digest)
    if suffix == ".xlsx":
        return _parse_xlsx(path, digest)
    if suffix == ".csv":
        return _parse_csv(path, digest, raw)
    if suffix == ".eml":
        return _parse_eml(path, digest, raw)
    if suffix == ".msg":
        return _parse_msg(path, digest)
    if suffix in {".txt", ".md"}:
        text = raw.decode("utf-8", errors="replace")
        return ParsedDocument(
            path=path,
            title=path.stem,
            text=text,
            mime_hint=suffix.lstrip("."),
            content_hash=digest,
            pages=[text],
        )
    raise ValueError(f"Unsupported file type: {path}")


def parse_directory(directory: Path) -> list[ParsedDocument]:
    docs: list[ParsedDocument] = []
    for path in sorted(directory.rglob("*")):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        # Skip macOS / junk
        if path.name.startswith(".") or path.name.startswith("__MACOSX"):
            continue
        docs.append(parse_file(path))
    return docs


def expand_zip(zip_path: Path, dest_dir: Path) -> list[Path]:
    """Extract a zip safely into dest_dir; return extracted file paths.

    Raises ValueError (and leaves nothing behind) when the archive holds more files or unpacks to
    more bytes than the settings allow. The size is counted while unpacking, not read from the
    archive's own header, so a ZIP bomb is stopped at the limit.
    """
    max_bytes = settings.zip_max_mb * 1024 * 1024
    dest_dir.mkdir(parents=True, exist_ok=True)
    extracted: list[Path] = []
    written = 0
    try:
        with zipfile.ZipFile(zip_path, "r") as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                name = Path(info.filename).name
                if not name or name.startswith("."):
                    continue
                if len(extracted) >= settings.zip_max_files:
                    raise ValueError(f"ZIP holds more than {settings.zip_max_files} files")
                # Zip-slip guard: only write basename into dest
                target = dest_dir / name
                # Avoid overwrite collisions
                if target.exists():
                    target = dest_dir / f"{target.stem}_{content_hash(info.filename.encode())[:6]}{target.suffix}"
                with zf.open(info) as src, open(target, "wb") as out:
                    while block := src.read(1024 * 1024):
                        written += len(block)
                        if written > max_bytes:
                            raise ValueError(f"ZIP unpacks to more than {settings.zip_max_mb} MB")
                        out.write(block)
                extracted.append(target)
    except Exception:
        shutil.rmtree(dest_dir, ignore_errors=True)
        raise
    return extracted


def _parse_pdf(path: Path, digest: str) -> ParsedDocument:
    doc = pymupdf.open(path)
    pages: list[str] = []
    for page in doc:
        pages.append(page.get_text("text"))
    page_count = doc.page_count
    doc.close()
    text = "\n\n".join(pages)
    warning = None
    if page_count > 0 and not text.strip():
        warning = (
            "PDF has no extractable text layer (likely a scan). "
            "OCR is out of MVP scope — re-export with text or OCR externally."
        )
    return ParsedDocument(
        path=path,
        title=path.stem,
        text=text,
        mime_hint="pdf",
        content_hash=digest,
        pages=pages,
        warning=warning,
    )


def _parse_docx(path: Path, digest: str) -> ParsedDocument:
    document = Document(path)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    return ParsedDocument(
        path=path,
        title=path.stem,
        text=text,
        mime_hint="docx",
        content_hash=digest,
        pages=[text],
    )


def _parse_pptx(path: Path, digest: str) -> ParsedDocument:
    prs = Presentation(path)
    slides: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[notes] {notes}")
        slides.append(f"[slide {i}]\n" + "\n".join(p for p in parts if p))
    text = "\n\n".join(slides)
    return ParsedDocument(
        path=path,
        title=path.stem,
        text=text,
        mime_hint="pptx",
        content_hash=digest,
        pages=slides,
    )


def _parse_xlsx(path: Path, digest: str) -> ParsedDocument:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets: list[str] = []
    for sheet in wb.worksheets:
        rows_out: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                rows_out.append("\t".join(cells))
        sheets.append(f"[sheet {sheet.title}]\n" + "\n".join(rows_out))
    wb.close()
    text = "\n\n".join(sheets)
    return ParsedDocument(
        path=path,
        title=path.stem,
        text=text,
        mime_hint="xlsx",
        content_hash=digest,
        pages=sheets,
    )


def _parse_csv(path: Path, digest: str, raw: bytes) -> ParsedDocument:
    text_in = raw.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text_in))
    rows = ["\t".join(row) for row in reader]
    text = "\n".join(rows)
    return ParsedDocument(
        path=path,
        title=path.stem,
        text=text,
        mime_hint="csv",
        content_hash=digest,
        pages=[text],
    )


def _parse_eml(path: Path, digest: str, raw: bytes) -> ParsedDocument:
    msg = BytesParser(policy=policy.default).parsebytes(raw)
    subject = msg.get("subject", path.stem)
    from_ = msg.get("from", "")
    to = msg.get("to", "")
    date = msg.get("date", "")
    body_parts: list[str] = []
    if msg.is_multipart():
        for part in msg.walk():
            ctype = part.get_content_type()
            if ctype == "text/plain":
                try:
                    body_parts.append(part.get_content())
                except Exception:
                    payload = part.get_payload(decode=True) or b""
                    body_parts.append(payload.decode("utf-8", errors="replace"))
            elif ctype == "text/html" and not body_parts:
                try:
                    html = part.get_content()
                except Exception:
                    payload = part.get_payload(decode=True) or b""
                    html = payload.decode("utf-8", errors="replace")
                # crude strip
                body_parts.append(_strip_html(html))
    else:
        try:
            body_parts.append(str(msg.get_content()))
        except Exception:
            payload = msg.get_payload(decode=True) or b""
            body_parts.append(payload.decode("utf-8", errors="replace"))

    header = f"Subject: {subject}\nFrom: {from_}\nTo: {to}\nDate: {date}\n"
    body = "\n".join(body_parts).strip()
    text = f"{header}\n{body}"
    return ParsedDocument(
        path=path,
        title=str(subject)[:200] or path.stem,
        text=text,
        mime_hint="eml",
        content_hash=digest,
        pages=[text],
    )


def _parse_msg(path: Path, digest: str) -> ParsedDocument:
    try:
        import extract_msg  # type: ignore
    except ImportError as exc:
        raise ValueError(".msg support requires extract-msg package") from exc

    msg = extract_msg.Message(str(path))
    try:
        subject = msg.subject or path.stem
        from_ = msg.sender or ""
        to = msg.to or ""
        date = str(msg.date or "")
        body = msg.body or msg.htmlBody or ""
        if isinstance(body, bytes):
            body = body.decode("utf-8", errors="replace")
        if "<html" in body.lower():
            body = _strip_html(body)
        header = f"Subject: {subject}\nFrom: {from_}\nTo: {to}\nDate: {date}\n"
        text = f"{header}\n{body}"
        return ParsedDocument(
            path=path,
            title=str(subject)[:200],
            text=text,
            mime_hint="msg",
            content_hash=digest,
            pages=[text],
        )
    finally:
        msg.close()


def _strip_html(html: str) -> str:
    import re

    text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", html)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", text).strip()
