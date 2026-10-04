"""Excel master-data export — consulting handoff format."""

from __future__ import annotations

import json
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from nlp_tda.config import settings
from nlp_tda.models import ArtifactRow, ProposedEntity

# Filter modes for review_status
FILTER_MODES = {
    "proposed": ("proposed",),
    "accepted": ("accepted",),
    "accepted_edited": ("accepted", "edited"),  # default handoff after review
    "all": ("proposed", "accepted", "edited", "rejected"),
}

DEFAULT_FILTER = "accepted_edited"

SHEET_SPECS: list[tuple[str, str, list[str]]] = [
    # (sheet_name, entity_type, payload-preferred columns after common cols)
    (
        "Client",
        "client",
        ["name", "aliases", "industry", "region", "status", "notes"],
    ),
    (
        "Engagement",
        "engagement",
        ["title", "client_name", "type", "phase", "start", "end", "status", "commercial_model"],
    ),
    (
        "Person",
        "person",
        ["name", "role", "org", "email", "engagement_titles"],
    ),
    (
        "Requirement",
        "requirement",
        ["statement", "priority", "status"],
    ),
    (
        "Finding",
        "finding",
        ["statement", "severity", "theme", "status"],
    ),
    (
        "Deliverable",
        "deliverable",
        ["name", "type", "due_date", "engagement_title"],
    ),
    (
        "Theme",
        "theme",
        ["label", "stability_score", "member_chunk_ids", "persistence_summary", "outlier"],
    ),
]

COMMON_COLS = [
    "id",
    "title",
    "review_status",
    "confidence",
    "source_artifact_id",
    "span_ref",
    "language",
    "run_id",
]

ARTIFACT_COLS = [
    "id",
    "title",
    "type",
    "path",
    "content_hash",
    "run_id",
    "text_preview",
]


def exports_dir() -> Path:
    path = Path(settings.exports_dir)
    path.mkdir(parents=True, exist_ok=True)
    return path


def resolve_statuses(filter_mode: str) -> tuple[str, ...]:
    if filter_mode not in FILTER_MODES:
        raise ValueError(f"Unknown filter_mode={filter_mode}; choose from {list(FILTER_MODES)}")
    return FILTER_MODES[filter_mode]


def count_by_entity(
    session: Session,
    *,
    filter_mode: str = DEFAULT_FILTER,
    run_id: str | None = None,
) -> dict[str, Any]:
    statuses = resolve_statuses(filter_mode)
    q = session.query(ProposedEntity).filter(ProposedEntity.review_status.in_(statuses))
    if run_id:
        q = q.filter(ProposedEntity.run_id == run_id)
    rows = q.all()
    counts: dict[str, int] = {spec[1]: 0 for spec in SHEET_SPECS}
    for row in rows:
        if row.entity_type in counts:
            counts[row.entity_type] += 1

    run_ids = {r.run_id for r in rows}
    aq = session.query(ArtifactRow)
    if run_id:
        aq = aq.filter(ArtifactRow.run_id == run_id)
        counts["artifact"] = aq.count()
    elif run_ids:
        counts["artifact"] = aq.filter(ArtifactRow.run_id.in_(run_ids)).count()
    elif filter_mode == "all":
        counts["artifact"] = session.query(ArtifactRow).count()
    else:
        counts["artifact"] = 0

    entity_rows = sum(v for k, v in counts.items() if k != "artifact")
    return {
        "filter_mode": filter_mode,
        "statuses": list(statuses),
        "run_id": run_id,
        "counts": counts,
        "entity_rows": entity_rows,
        "total_entities": entity_rows,
        "total_with_artifacts": entity_rows + counts["artifact"],
    }


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, bool):
        return value
    return value


def _write_sheet(ws, headers: list[str], data_rows: list[list[Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in data_rows:
        ws.append([_cell(v) for v in row])
    for idx, _ in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        maxlen = len(str(headers[idx - 1]))
        for row in data_rows[:50]:
            if idx - 1 < len(row):
                maxlen = max(maxlen, min(len(str(_cell(row[idx - 1]))), 48))
        ws.column_dimensions[letter].width = max(12, maxlen + 2)


def build_workbook(
    session: Session,
    *,
    filter_mode: str = DEFAULT_FILTER,
    run_id: str | None = None,
) -> tuple[Workbook, dict[str, Any]]:
    statuses = resolve_statuses(filter_mode)
    q = session.query(ProposedEntity).filter(ProposedEntity.review_status.in_(statuses))
    if run_id:
        q = q.filter(ProposedEntity.run_id == run_id)
    entities = q.order_by(ProposedEntity.entity_type, ProposedEntity.id).all()

    by_type: dict[str, list[ProposedEntity]] = {}
    for ent in entities:
        by_type.setdefault(ent.entity_type, []).append(ent)

    run_ids = {e.run_id for e in entities}
    aq = session.query(ArtifactRow)
    if run_id:
        artifacts = aq.filter(ArtifactRow.run_id == run_id).all()
    elif run_ids:
        artifacts = aq.filter(ArtifactRow.run_id.in_(run_ids)).all()
    elif filter_mode == "all":
        artifacts = aq.all()
    else:
        artifacts = []

    wb = Workbook()
    # Summary first
    summary = wb.active
    summary.title = "Summary"
    summary.append(["NLP_TDA master-data export"])
    summary.append(["generated_at_utc", datetime.now(timezone.utc).isoformat()])
    summary.append(["filter_mode", filter_mode])
    summary.append(["statuses", ", ".join(statuses)])
    summary.append(["run_id", run_id or "(all matching)"])
    summary.append([])
    summary.append(["sheet", "rows"])
    counts: dict[str, int] = {}

    for sheet_name, entity_type, payload_cols in SHEET_SPECS:
        ws = wb.create_sheet(sheet_name)
        headers = COMMON_COLS + [c for c in payload_cols if c not in COMMON_COLS]
        data_rows: list[list[Any]] = []
        for ent in by_type.get(entity_type, []):
            payload = dict(ent.payload or {})
            row = [
                ent.id,
                ent.title,
                ent.review_status,
                ent.confidence,
                ent.source_artifact_id or payload.get("source_artifact_id") or "",
                ent.span_ref or payload.get("span_ref") or "",
                ent.language or payload.get("language") or "",
                ent.run_id,
            ]
            for col in payload_cols:
                if col in COMMON_COLS:
                    continue
                row.append(payload.get(col, ""))
            data_rows.append(row)
        _write_sheet(ws, headers, data_rows)
        counts[entity_type] = len(data_rows)
        summary.append([sheet_name, len(data_rows)])

    # Artifact sheet from ingest store
    art_ws = wb.create_sheet("Artifact")
    art_rows = [
        [
            a.id,
            a.title,
            a.mime_hint,
            a.path,
            a.content_hash,
            a.run_id,
            (a.text_preview or "")[:500],
        ]
        for a in artifacts
    ]
    _write_sheet(art_ws, ARTIFACT_COLS, art_rows)
    counts["artifact"] = len(art_rows)
    summary.append(["Artifact", len(art_rows)])
    summary.append([])
    summary.append(["total_entity_rows", sum(counts[k] for k in counts if k != "artifact")])
    summary.append(["total_with_artifacts", sum(counts.values())])

    meta = {
        "filter_mode": filter_mode,
        "statuses": list(statuses),
        "run_id": run_id,
        "counts": counts,
        "total_entities": sum(counts[k] for k in counts if k != "artifact"),
        "total_with_artifacts": sum(counts.values()),
    }
    return wb, meta


def export_excel(
    session: Session,
    *,
    filter_mode: str = DEFAULT_FILTER,
    run_id: str | None = None,
) -> dict[str, Any]:
    wb, meta = build_workbook(session, filter_mode=filter_mode, run_id=run_id)
    file_id = uuid.uuid4().hex[:12]
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    filename = f"nlp_tda_master_{filter_mode}_{stamp}_{file_id}.xlsx"
    path = exports_dir() / filename
    wb.save(path)
    return {
        "file_id": file_id,
        "filename": filename,
        "path": str(path),
        **meta,
    }


def find_export(file_id: str) -> Path | None:
    matches = list(exports_dir().glob(f"*_{file_id}.xlsx"))
    return matches[0] if matches else None
