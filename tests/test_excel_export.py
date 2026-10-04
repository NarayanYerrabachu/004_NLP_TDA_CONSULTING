from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from nlp_tda.config import settings
from nlp_tda.db import get_session, reset_db_for_tests
from nlp_tda.export.excel_export import export_excel
from nlp_tda.models import ProposedEntity, ReviewStatus
from nlp_tda.pipeline import run_pipeline

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "synthetic_engagement"


def test_excel_export_sheets_and_rows(tmp_path):
    settings.data_dir = tmp_path / "data"
    settings.sqlite_path = tmp_path / "data" / "master.db"
    settings.chroma_path = tmp_path / "data" / "chroma"
    settings.uploads_dir = tmp_path / "data" / "uploads"
    settings.exports_dir = tmp_path / "data" / "exports"
    settings.project_root = tmp_path
    settings.force_mock_llm = True
    settings.use_hash_embeddings = True
    reset_db_for_tests()

    result = run_pipeline(FIXTURES, force_hash_embeddings=True)
    session = get_session()
    try:
        # Accept a couple so accepted_edited is non-empty; also export proposed
        rows = session.query(ProposedEntity).limit(3).all()
        for row in rows:
            row.review_status = ReviewStatus.accepted.value
        session.commit()

        proposed = export_excel(session, filter_mode="proposed", run_id=result["run_id"])
        accepted = export_excel(session, filter_mode="accepted", run_id=result["run_id"])
        handoff = export_excel(session, filter_mode="accepted_edited", run_id=result["run_id"])
        everything = export_excel(session, filter_mode="all", run_id=result["run_id"])
    finally:
        session.close()

    assert proposed["counts"]["theme"] >= 1
    assert accepted["total_entities"] >= 3
    assert handoff["total_entities"] >= 3
    assert everything["total_with_artifacts"] > everything["total_entities"]

    wb = load_workbook(accepted["path"])
    expected = {"Summary", "Client", "Engagement", "Person", "Requirement", "Finding", "Deliverable", "Theme", "Artifact"}
    assert expected <= set(wb.sheetnames)
    # Accepted sheets should have header + at least some entity rows across workbook
    assert sum(max(ws.max_row - 1, 0) for ws in wb.worksheets if ws.title != "Summary") >= 3
